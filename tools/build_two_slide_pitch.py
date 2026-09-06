#!/usr/bin/env python3
"""Create an editable two-slide EPIC 8 pitch deck.

This is intentionally made from native PowerPoint shapes rather than a rendered
image so the two member names and the GitHub / demo URLs remain easy to replace.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).resolve().parents[1] / "deliverables" / "AI_Smart_Lending_Two_Slide_Pitch.pptx"
W, H = Inches(13.333), Inches(7.5)

# Visual language taken from Lending_Hub_EPIC_8_v2.pptx.
NAVY = RGBColor(0x0A, 0x3A, 0x68)
BLUE = RGBColor(0x21, 0x66, 0xA3)
PALE_BLUE = RGBColor(0xEE, 0xF4, 0xFA)
GREEN = RGBColor(0x06, 0x76, 0x47)
PALE_GREEN = RGBColor(0xE8, 0xF5, 0xEF)
AMBER = RGBColor(0xB5, 0x47, 0x08)
PALE_AMBER = RGBColor(0xFD, 0xF3, 0xEA)
INK = RGBColor(0x10, 0x18, 0x28)
MUTED = RGBColor(0x58, 0x66, 0x7A)
RULE = RGBColor(0xDB, 0xE3, 0xEC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BODY, DISPLAY, MONO = "Aptos", "Georgia", "Consolas"


def add_text(slide, text, x, y, w, h, *, size=12, color=INK, bold=False,
             font=BODY, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(x, y, w, h)
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    paragraph.text = text
    for run in paragraph.runs:
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return shape


def rect(slide, x, y, w, h, fill=WHITE, line=RULE, radius=False):
    kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(0.7)
    return shape


def line(slide, x1, y1, x2, y2, color=RULE, width=1.0, arrow=False):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    if arrow:
        shape.line.end_arrowhead = True
    return shape


def label(slide, text, x, y, w, *, color=BLUE):
    add_text(slide, text.upper(), x, y, w, Inches(0.18), size=8.5, color=color, bold=True, font=MONO)


def footer(slide, number):
    line(slide, Inches(0.48), Inches(6.92), Inches(12.85), Inches(6.92), RULE, 0.7)
    add_text(slide, "TEAM: <MEMBER 1 NAME>  •  <MEMBER 2 NAME>", Inches(0.5), Inches(7.05), Inches(4.2), Inches(0.18), size=8.3, color=MUTED, font=MONO)
    add_text(slide, "GITHUB: <PASTE REPOSITORY URL>", Inches(4.75), Inches(7.05), Inches(3.55), Inches(0.18), size=8.3, color=BLUE, font=MONO)
    add_text(slide, "YOUTUBE DEMO: <PASTE VIDEO URL>", Inches(8.35), Inches(7.05), Inches(3.85), Inches(0.18), size=8.3, color=BLUE, font=MONO)
    add_text(slide, f"0{number}", Inches(12.35), Inches(7.02), Inches(0.45), Inches(0.2), size=9, color=MUTED, font=MONO, align=PP_ALIGN.RIGHT)


def icon_circle(slide, value, x, y, *, fill=BLUE):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, x, y, Inches(0.34), Inches(0.34))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    add_text(slide, value, x, y + Inches(0.045), Inches(0.34), Inches(0.18), size=9, color=WHITE, bold=True, font=MONO, align=PP_ALIGN.CENTER)


def problem_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = WHITE
    rect(slide, Inches(0), Inches(0), Inches(0.18), H, NAVY, NAVY)
    label(slide, "TVS Credit EPIC 8.0  ·  AI-powered smart lending decision hub", Inches(0.55), Inches(0.38), Inches(8.5))
    add_text(slide, "Agricultural lending has an information gap —\nnot just a credit-score gap.", Inches(0.55), Inches(0.68), Inches(8.3), Inches(0.86), size=25, color=NAVY, bold=True, font=DISPLAY)
    add_text(slide, "Farmers may have limited formal credit history, while lenders lack timely crop, cash-flow and fraud context. The result is slower, less inclusive lending and risk that is discovered after the borrower is already in trouble.", Inches(0.57), Inches(1.62), Inches(8.05), Inches(0.55), size=11.5, color=MUTED)

    # Central causal flow: the core problem story.
    label(slide, "THE LENDING GAP", Inches(0.57), Inches(2.33), Inches(2.2), color=AMBER)
    stages = [
        ("01", "Thin borrower evidence", "Limited bureau history, seasonal income and fragmented records"),
        ("02", "Uncertain lending decision", "Slow manual review or a generic score misses the farm context"),
        ("03", "Late risk visibility", "Fraud, crop stress and repayment deterioration appear too late"),
    ]
    xs = [Inches(0.57), Inches(3.55), Inches(6.53)]
    for idx, ((num, title, body), x) in enumerate(zip(stages, xs)):
        rect(slide, x, Inches(2.63), Inches(2.56), Inches(1.4), PALE_AMBER, RGBColor(0xF0, 0xD7, 0xB8), True)
        icon_circle(slide, num, x + Inches(0.16), Inches(2.82), fill=AMBER)
        add_text(slide, title, x + Inches(0.62), Inches(2.79), Inches(1.74), Inches(0.28), size=11.5, color=INK, bold=True)
        add_text(slide, body, x + Inches(0.17), Inches(3.2), Inches(2.18), Inches(0.55), size=9.3, color=MUTED)
        if idx < 2:
            line(slide, x + Inches(2.58), Inches(3.33), x + Inches(2.92), Inches(3.33), AMBER, 1.4, True)

    # The two stakeholder consequences.
    rect(slide, Inches(9.52), Inches(0.5), Inches(3.18), Inches(3.53), NAVY, NAVY, True)
    label(slide, "WHAT THIS COSTS", Inches(9.82), Inches(0.78), Inches(2.2), color=RGBColor(0xAD, 0xC9, 0xE5))
    add_text(slide, "For farmers", Inches(9.82), Inches(1.1), Inches(2), Inches(0.25), size=12, color=WHITE, bold=True)
    add_text(slide, "• Delayed or declined credit\n• One-size-fits-all offers\n• No early support when stress starts", Inches(9.82), Inches(1.42), Inches(2.52), Inches(0.85), size=10, color=RGBColor(0xE6, 0xF0, 0xF8))
    line(slide, Inches(9.82), Inches(2.43), Inches(12.35), Inches(2.43), RGBColor(0x4D, 0x78, 0xA5), 0.8)
    add_text(slide, "For lenders", Inches(9.82), Inches(2.63), Inches(2), Inches(0.25), size=12, color=WHITE, bold=True)
    add_text(slide, "• Higher manual effort\n• Hidden fraud and concentration risk\n• Intervention arrives after default", Inches(9.82), Inches(2.95), Inches(2.52), Inches(0.85), size=10, color=RGBColor(0xE6, 0xF0, 0xF8))

    # Grounding / trust differentiator.
    rect(slide, Inches(0.57), Inches(4.5), Inches(12.13), Inches(1.92), PALE_BLUE, RGBColor(0xC4, 0xD8, 0xEA), True)
    label(slide, "THE DESIGN PRINCIPLE", Inches(0.84), Inches(4.76), Inches(2.5))
    add_text(slide, "A lending model is easy to build. Knowing whether you are allowed to trust it is the hard part.", Inches(0.84), Inches(5.04), Inches(5.65), Inches(0.52), size=16, color=NAVY, bold=True, font=DISPLAY)
    add_text(slide, "Every output must be grounded in one of three sources — otherwise the platform refuses to invent a number.", Inches(0.84), Inches(5.7), Inches(5.7), Inches(0.32), size=10.3, color=MUTED)
    sources = [("[SPEC]", "written rule / formula"), ("[DATA]", "rerunnable computation"), ("[POLICY]", "named owner approval")]
    for i, (tag, desc) in enumerate(sources):
        x = Inches(6.85 + i * 1.87)
        rect(slide, x, Inches(5.0), Inches(1.62), Inches(0.95), WHITE, RULE, True)
        add_text(slide, tag, x + Inches(0.12), Inches(5.18), Inches(1.36), Inches(0.2), size=9.4, color=BLUE if i < 2 else AMBER, bold=True, font=MONO, align=PP_ALIGN.CENTER)
        add_text(slide, desc, x + Inches(0.11), Inches(5.53), Inches(1.4), Inches(0.18), size=7.8, color=MUTED, align=PP_ALIGN.CENTER)
    footer(slide, 1)


def solution_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = WHITE
    rect(slide, Inches(0), Inches(0), W, Inches(0.16), NAVY, NAVY)
    label(slide, "THE SOLUTION  ·  ONE GOVERNED DECISION FLOW", Inches(0.55), Inches(0.38), Inches(5.2))
    add_text(slide, "From fragmented signals to earlier, explainable action.", Inches(0.55), Inches(0.69), Inches(8.4), Inches(0.42), size=24, color=NAVY, bold=True, font=DISPLAY)
    add_text(slide, "AI-Powered Smart Lending Decision Hub unifies agricultural intelligence, risk models and human-ready evidence through one decision orchestrator.", Inches(0.57), Inches(1.22), Inches(9.25), Inches(0.3), size=11.3, color=MUTED)

    # Inputs → gateway → actions diagram.
    label(slide, "1. EVIDENCE", Inches(0.57), Inches(1.8), Inches(1.6))
    inputs = [("SAT", "Satellite + crop"), ("WX", "Weather + location"), ("CR", "Credit + cash flow"), ("FR", "Device + documents")]
    for i, (code, name) in enumerate(inputs):
        y = Inches(2.12 + i * 0.56)
        rect(slide, Inches(0.57), y, Inches(2.1), Inches(0.42), PALE_BLUE, RGBColor(0xC4, 0xD8, 0xEA), True)
        add_text(slide, code, Inches(0.73), y + Inches(0.12), Inches(0.34), Inches(0.14), size=8, color=BLUE, bold=True, font=MONO)
        add_text(slide, name, Inches(1.13), y + Inches(0.105), Inches(1.3), Inches(0.16), size=9.5, color=INK, bold=True)
        line(slide, Inches(2.7), y + Inches(0.21), Inches(3.12), y + Inches(0.21), BLUE, 1.0, True)

    rect(slide, Inches(3.17), Inches(2.06), Inches(2.45), Inches(2.52), NAVY, NAVY, True)
    add_text(slide, "DECISION\nORCHESTRATOR", Inches(3.48), Inches(2.4), Inches(1.84), Inches(0.55), size=15, color=WHITE, bold=True, font=DISPLAY, align=PP_ALIGN.CENTER)
    add_text(slide, "One gateway\nAttribution on every\nmodel-derived output", Inches(3.55), Inches(3.25), Inches(1.7), Inches(0.65), size=9.4, color=RGBColor(0xD6, 0xE4, 0xF2), align=PP_ALIGN.CENTER)
    add_text(slide, "evidence → decision → audit", Inches(3.35), Inches(4.16), Inches(2.08), Inches(0.16), size=7.4, color=RGBColor(0xAD, 0xC9, 0xE5), font=MONO, align=PP_ALIGN.CENTER)

    line(slide, Inches(5.65), Inches(3.32), Inches(6.06), Inches(3.32), GREEN, 1.4, True)
    label(slide, "2. CONNECTED INTELLIGENCE LAYER · ALL 8 MODULES", Inches(6.15), Inches(1.8), Inches(4.0), color=GREEN)
    engines = [
        ("Agri intelligence", "satellite · weather · crop"), ("Credit scoring", "risk + explanations"),
        ("Fraud detection", "rules · anomaly · graph"), ("Default prediction", "PD · LGD · EAD"),
        ("Loan recommendation", "product · price · tenor"), ("Early warning", "detect · route · act"),
        ("GenAI assistant", "grounded guidance"), ("Risk dashboards", "portfolio · drill-down"),
    ]
    for i, (title, detail) in enumerate(engines):
        col, row = i % 2, i // 2
        x, y = Inches(6.15 + col * 2.15), Inches(2.12 + row * 0.72)
        rect(slide, x, y, Inches(1.92), Inches(0.56), PALE_GREEN, RGBColor(0xB9, 0xDF, 0xCC), True)
        add_text(slide, title, x + Inches(0.12), y + Inches(0.11), Inches(1.68), Inches(0.16), size=8.8, color=INK, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, detail, x + Inches(0.12), y + Inches(0.31), Inches(1.68), Inches(0.13), size=7.1, color=GREEN, align=PP_ALIGN.CENTER)

    line(slide, Inches(10.3), Inches(3.32), Inches(10.69), Inches(3.32), GREEN, 1.4, True)
    label(slide, "3. OUTCOMES", Inches(10.77), Inches(1.8), Inches(1.7), color=AMBER)
    outcomes = [("FASTER", "informed approval"), ("SAFER", "fraud-aware review"), ("EARLIER", "risk intervention")]
    for i, (head, sub) in enumerate(outcomes):
        y = Inches(2.15 + i * 0.77)
        rect(slide, Inches(10.74), y, Inches(1.95), Inches(0.56), PALE_AMBER, RGBColor(0xF0, 0xD7, 0xB8), True)
        add_text(slide, head, Inches(10.88), y + Inches(0.1), Inches(1.65), Inches(0.17), size=9, color=AMBER, bold=True, font=MONO, align=PP_ALIGN.CENTER)
        add_text(slide, sub, Inches(10.88), y + Inches(0.32), Inches(1.65), Inches(0.13), size=7.8, color=MUTED, align=PP_ALIGN.CENTER)

    # Evidence and presentation close.
    rect(slide, Inches(0.57), Inches(5.08), Inches(12.12), Inches(1.32), PALE_BLUE, RGBColor(0xC4, 0xD8, 0xEA), True)
    label(slide, "BUILT, MEASURED, AND DESIGNED FOR GOVERNANCE", Inches(0.84), Inches(5.31), Inches(4.5))
    metrics = [("19", "models across 4 families"), ("590K", "fraud transactions · 0.9564 AUC"), ("307K", "loan applications for credit scoring"), ("365d", "median EWS lead time at p95")]
    for i, (value, caption) in enumerate(metrics):
        x = Inches(0.85 + i * 2.9)
        add_text(slide, value, x, Inches(5.67), Inches(0.85), Inches(0.34), size=20, color=NAVY if i != 1 else GREEN, bold=True, font=DISPLAY)
        add_text(slide, caption, x + Inches(0.92), Inches(5.73), Inches(1.72), Inches(0.28), size=8.4, color=MUTED)
    footer(slide, 2)


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    problem_slide(prs)
    solution_slide(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
